#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
wb_stocks — проверка остатков товаров на Wildberries по списку артикулов.

Читает артикулы из xlsx/csv/txt, опрашивает публичный API карточек WB
и складывает результат в отчёт Остатки_WB_<дата>.xlsx.

Главное отличие от прежней версии: запрос к WB больше не «прибит гвоздями».
Перед основным прогоном выполняется подбор рабочей комбинации
(эндпоинт + dest + профиль браузера), сессия прогревается на www.wildberries.ru,
а батчи при отказе дробятся. Это и лечит сплошной 403 Forbidden.
"""

from __future__ import annotations

import json
import os
import random
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import parse_qs, quote

# --------------------------------------------------------------------------
# HTTP-движок
# --------------------------------------------------------------------------

try:
    from curl_cffi import requests as http_lib

    HAS_CURL_CFFI = True
except ImportError:  # pragma: no cover - запасной путь, если curl_cffi не собрался
    import requests as http_lib  # type: ignore[no-redef]

    HAS_CURL_CFFI = False


# --------------------------------------------------------------------------
# Константы
# --------------------------------------------------------------------------

WB_MAIN = "https://www.wildberries.ru/"
GEO_URL = "https://user-geo-data.wildberries.ru/get-geo-info"
STATS_URL = "https://statistics-api.wildberries.ru/api/v1/supplier/stocks"

# Кандидаты-эндпоинты карточек. Перебираются сверху вниз при подборе.
CARD_ENDPOINTS: list[tuple[str, int]] = [
    ("https://card.wb.ru/cards/v2/detail", 2),
    ("https://u-card.wb.ru/cards/v4/detail", 4),
    ("https://card.wb.ru/cards/v4/detail", 4),
    ("https://card.wb.ru/cards/detail", 1),
]

# Коды региона доставки. Без валидного dest WB отвечает 403/400.
# Запасные коды региона на случай, если гео-сервис недоступен.
# Формат WB сменил: актуальные коды — длинные положительные числа,
# старые отрицательные оставлены последними как последняя надежда.
DEST_CANDIDATES: list[str] = [
    "1259570991",  # Москва, действующий формат
    "-1257786",    # Москва, старый формат
    "-1029256",    # Санкт-Петербург, старый формат
]

IMPERSONATE_PROFILES: list[str] = [
    "chrome131",
    "chrome124",
    "chrome120",
    "chrome116",
    "edge101",
    "safari17_0",
]

FALLBACK_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
)

# WB режет крупные пачки. 100 штук в запросе — одна из причин прежнего 403.
DEFAULT_BATCH = 50
MIN_BATCH = 5
REQUEST_TIMEOUT = 30
MAX_ATTEMPTS = 4          # попыток на один запрос до дробления батча
MAX_DEAD_STREAK = 40      # подряд неудачных запросов, после которых не дробим
MAX_PROBE_REQUESTS = 36   # предел подбора канала, чтобы не долбить WB впустую
MAX_PROBE_ROUNDS = 3      # столько профилей браузера перебираем при подборе

# Куда сходить, чтобы узнать свой внешний IP и страну (проверка прокси).
IP_CHECK_URLS = [
    "https://ipinfo.io/json",
    "https://api.myip.com",
    "https://api.ipify.org?format=json",
]

PROXY_SCHEMES = ("http", "https", "socks5", "socks5h", "socks4", "socks4a")

STATUS_ON_SALE = "в продаже"
STATUS_ZERO = "нулевой остаток"
STATUS_ABSENT = "нет на сайте"
STATUS_ERROR = "ошибка запроса"


# --------------------------------------------------------------------------
# Утилиты окружения
# --------------------------------------------------------------------------

def setup_console() -> None:
    """UTF-8 в консоли Windows, иначе кириллица в exe превращается в кашу."""
    for stream in (sys.stdout, sys.stderr):
        try:
            if hasattr(stream, "reconfigure"):
                stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def base_dir() -> Path:
    """Папка рядом с exe (в собранном виде) или со скриптом."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


class Log:
    """Печать в консоль + полный протокол в файл рядом с программой."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self._fh = None
        try:
            self._fh = path.open("w", encoding="utf-8")
        except Exception:
            self._fh = None

    def __call__(self, text: str = "", to_console: bool = True) -> None:
        if to_console:
            print(text, flush=True)
        if self._fh:
            try:
                stamp = datetime.now().strftime("%H:%M:%S")
                self._fh.write(f"[{stamp}] {text}\n")
                self._fh.flush()
            except Exception:
                pass

    def detail(self, text: str) -> None:
        """Только в файл — подробности, которыми не надо засорять консоль."""
        self(text, to_console=False)

    def close(self) -> None:
        if self._fh:
            try:
                self._fh.close()
            except Exception:
                pass


# --------------------------------------------------------------------------
# Прокси
# --------------------------------------------------------------------------

def normalize_proxy(raw: str) -> str:
    """Приводит строку прокси к виду scheme://user:pass@host:port.

    Продавцы отдают адреса в разном виде, поэтому понимаем все ходовые:
        host:port
        host:port:user:pass          (proxy6, proxy-seller и подобные)
        user:pass@host:port
        socks5://user:pass@host:port
    """
    raw = (raw or "").strip()
    if not raw or raw.startswith("#"):
        return ""

    scheme = "http"
    if "://" in raw:
        scheme, _, raw = raw.partition("://")
        scheme = scheme.lower()
        if scheme not in PROXY_SCHEMES:
            raise ValueError(f"неизвестный протокол прокси: {scheme}")
        # socks5h заставляет резолвить домены на стороне прокси:
        # иначе DNS идёт от нас и легко упирается в местного провайдера.
        if scheme == "socks5":
            scheme = "socks5h"

    user = password = ""
    if "@" in raw:
        credentials, _, raw = raw.rpartition("@")
        user, _, password = credentials.partition(":")

    parts = raw.split(":")
    if len(parts) == 4 and not user:
        # host:port:user:pass
        host, port, user, password = parts
    elif len(parts) == 2:
        host, port = parts
    else:
        raise ValueError(f"не разобрал адрес прокси: {raw}")

    host = host.strip()
    port = port.strip()
    if not host or not port.isdigit():
        raise ValueError(f"не разобрал адрес прокси: {raw}")

    if user:
        # Пароли часто содержат спецсимволы — их надо экранировать,
        # иначе ломается разбор URL внутри curl.
        auth = f"{quote(user, safe='')}:{quote(password, safe='')}@"
    else:
        auth = ""
    return f"{scheme}://{auth}{host}:{port}"


def mask_proxy(url: str) -> str:
    """Прячет логин и пароль, чтобы не светить их в консоли и логе."""
    if not url:
        return "без прокси"
    scheme, _, rest = url.partition("://")
    if "@" in rest:
        rest = rest.rpartition("@")[2]
    return f"{scheme}://{rest}"


# --------------------------------------------------------------------------
# Конфигурация
# --------------------------------------------------------------------------

@dataclass
class Config:
    token: str = ""
    dest: str = ""
    batch_size: int = DEFAULT_BATCH
    delay: float = 0.4
    input_file: str = ""
    proxies: list[str] = field(default_factory=list)

    @classmethod
    def load(cls, folder: Path, log: Log) -> "Config":
        cfg = cls()
        raw_proxies: list[str] = []

        path = folder / "wb_config.json"
        if path.exists():
            try:
                raw = json.loads(path.read_text(encoding="utf-8-sig"))
                cfg.token = str(raw.get("token", "") or "")
                cfg.dest = str(raw.get("dest", "") or "")
                cfg.batch_size = int(raw.get("batch_size", DEFAULT_BATCH) or DEFAULT_BATCH)
                cfg.delay = float(raw.get("delay", 0.4) or 0.4)
                cfg.input_file = str(raw.get("input_file", "") or "")
                if raw.get("proxy"):
                    raw_proxies.append(str(raw["proxy"]))
                if isinstance(raw.get("proxies"), list):
                    raw_proxies.extend(str(item) for item in raw["proxies"])
                log(f"Настройки: {path.name}")
            except Exception as exc:
                log(f"! wb_config.json не прочитан ({exc}), беру настройки по умолчанию")

        proxy_file = folder / "proxy.txt"
        if proxy_file.exists():
            try:
                raw_proxies.extend(proxy_file.read_text(encoding="utf-8-sig").splitlines())
            except Exception as exc:
                log(f"! proxy.txt не прочитан: {exc}")

        env_proxy = os.environ.get("WB_PROXY", "").strip()
        if env_proxy:
            raw_proxies.append(env_proxy)

        for item in raw_proxies:
            try:
                normalized = normalize_proxy(item)
            except ValueError as exc:
                log(f"! прокси пропущен — {exc}")
                continue
            if normalized and normalized not in cfg.proxies:
                cfg.proxies.append(normalized)

        token_file = folder / "token.txt"
        if not cfg.token and token_file.exists():
            try:
                cfg.token = token_file.read_text(encoding="utf-8-sig").strip()
            except Exception:
                pass
        if not cfg.token:
            cfg.token = os.environ.get("WB_API_TOKEN", "").strip()

        cfg.batch_size = max(1, min(cfg.batch_size, 100))
        cfg.delay = max(0.0, min(cfg.delay, 10.0))
        return cfg


# --------------------------------------------------------------------------
# Чтение артикулов
# --------------------------------------------------------------------------

NM_RE = re.compile(r"\b(\d{5,12})\b")


# Служебные файлы программы — их нельзя принять за список артикулов.
SERVICE_FILES = {
    "token.txt",
    "proxy.txt",
    "requirements.txt",
    "wb_config.json",
    "wb_stocks.log",
    "readme.txt",
    "readme.md",
}


def find_input_file(folder: Path, preferred: str = "") -> Path | None:
    if preferred:
        candidate = Path(preferred)
        if not candidate.is_absolute():
            candidate = folder / candidate
        if candidate.exists():
            return candidate

    default = folder / "шаблон_артикулы.xlsx"
    if default.exists():
        return default

    for pattern in ("*.xlsx", "*.csv", "*.txt"):
        for path in sorted(folder.glob(pattern)):
            name = path.name
            if name.startswith("~$") or name.startswith("Остатки_WB"):
                continue
            if name.lower() in SERVICE_FILES:
                continue
            return path
    return None


def extract_nm(value: Any) -> int | None:
    """Достаёт артикул из числа, строки или ссылки на карточку."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = int(value)
        return number if 10_000 <= number <= 999_999_999_999 else None

    text = str(value).strip()
    if not text:
        return None

    link = re.search(r"/catalog/(\d{5,12})/", text)
    if link:
        return int(link.group(1))

    match = NM_RE.search(text.replace(" ", ""))
    return int(match.group(1)) if match else None


def read_articles(path: Path) -> list[int]:
    suffix = path.suffix.lower()
    values: list[Any] = []

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values.extend(row)
        finally:
            book.close()
    else:
        encodings = ("utf-8-sig", "cp1251", "utf-8")
        text = ""
        for encoding in encodings:
            try:
                text = path.read_text(encoding=encoding)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        values.extend(re.split(r"[\r\n;,\t]+", text))

    seen: set[int] = set()
    articles: list[int] = []
    for value in values:
        nm = extract_nm(value)
        if nm is not None and nm not in seen:
            seen.add(nm)
            articles.append(nm)
    return articles


# --------------------------------------------------------------------------
# Разбор ответа WB
# --------------------------------------------------------------------------

def parse_products(payload: Any) -> list[dict]:
    """v2 кладёт товары в data.products, v4 — в products верхнего уровня."""
    if not isinstance(payload, dict):
        return []
    data = payload.get("data")
    if isinstance(data, dict) and isinstance(data.get("products"), list):
        return [p for p in data["products"] if isinstance(p, dict)]
    if isinstance(payload.get("products"), list):
        return [p for p in payload["products"] if isinstance(p, dict)]
    return []


def product_stock(product: dict) -> int:
    total = 0
    counted = False
    for size in product.get("sizes") or []:
        if not isinstance(size, dict):
            continue
        for stock in size.get("stocks") or []:
            if isinstance(stock, dict) and isinstance(stock.get("qty"), (int, float)):
                total += int(stock["qty"])
                counted = True
    if not counted:
        # volume сюда не берём — это объём упаковки, а не остаток.
        for key in ("totalQuantity", "qty"):
            value = product.get(key)
            if isinstance(value, (int, float)):
                return int(value)
    return total


def product_price(product: dict) -> float | None:
    for size in product.get("sizes") or []:
        if not isinstance(size, dict):
            continue
        price = size.get("price")
        if isinstance(price, dict):
            for key in ("total", "product", "basic"):
                value = price.get(key)
                if isinstance(value, (int, float)) and value > 0:
                    return round(value / 100, 2)
    for key in ("salePriceU", "priceU"):
        value = product.get(key)
        if isinstance(value, (int, float)) and value > 0:
            return round(value / 100, 2)
    return None


# --------------------------------------------------------------------------
# Клиент WB
# --------------------------------------------------------------------------

@dataclass
class Attempt:
    """Что именно пробовали — для диагностики, когда всё легло."""
    endpoint: str
    dest: str
    profile: str
    result: str


class WBClient:
    def __init__(self, log: Log, delay: float = 0.4, dest_hint: str = "",
                 proxies: Sequence[str] = ()) -> None:
        self.log = log
        self.delay = delay
        self.dest_hint = dest_hint
        self.proxies = list(proxies)
        self.proxy_index = 0
        self.endpoint: str = CARD_ENDPOINTS[0][0]
        self.version: int = CARD_ENDPOINTS[0][1]
        self.dest: str = DEST_CANDIDATES[0]
        self.profile_index = 0
        self.session = None
        self.dead_streak = 0
        self.batch_size = DEFAULT_BATCH
        self.attempts: list[Attempt] = []
        # Результат прогрева витрины: без его кук карточки отдают 403,
        # поэтому при разборе полётов это первое, что надо знать.
        self.warm_status = "не выполнялся"
        self.warm_cookies = ""

    # -- сессия ------------------------------------------------------------

    @property
    def profile(self) -> str:
        return IMPERSONATE_PROFILES[self.profile_index % len(IMPERSONATE_PROFILES)]

    @property
    def proxy(self) -> str:
        if not self.proxies:
            return ""
        return self.proxies[self.proxy_index % len(self.proxies)]

    def proxy_map(self) -> dict[str, str] | None:
        proxy = self.proxy
        return {"http": proxy, "https": proxy} if proxy else None

    def report_ip(self) -> bool:
        """Печатает внешний IP и страну текущего канала. False — не ответил."""
        for url in IP_CHECK_URLS:
            payload, error = self.raw_request(url, {})
            if error or not isinstance(payload, dict):
                self.log.detail(f"проверка IP через {url}: {error or 'нет данных'}")
                continue

            ip = payload.get("ip") or payload.get("query") or "?"
            country = (payload.get("country") or payload.get("cc") or "").upper()
            where = mask_proxy(self.proxy) if self.proxy else "прямое подключение"
            suffix = f", страна {country}" if country else ""
            self.log(f"Внешний IP: {ip}{suffix} ({where})")

            if country and country != "RU":
                self.log("! IP не российский — WB такие адреса режет целиком.")
                self.log("  Нужен прокси с российским IP, иначе 403 никуда не денется.")
            return True

        self.log.detail(f"внешний IP не определён ({mask_proxy(self.proxy)})")
        return False

    def check_proxy(self) -> None:
        """Ищет живой прокси и показывает, с какого IP реально пойдут запросы.

        Для WB страна IP — вопрос номер один, и узнать ответ здесь дешевле,
        чем ловить 403 на каждом запросе. Мёртвые адреса пропускаем сразу,
        чтобы подбор канала стартовал уже с рабочего прокси.
        """
        if not self.proxies:
            self.new_session(warm_up=False)
            self.report_ip()
            return

        # Длинные списки целиком не проверяем — это долгий старт без пользы.
        for _ in range(min(len(self.proxies), 5)):
            self.new_session(warm_up=False)
            if self.report_ip():
                return
            self.log(f"! Прокси {mask_proxy(self.proxy)} не отвечает — беру следующий")
            self.proxy_index += 1

        self.log("! Ни один из проверенных прокси не ответил.")
        self.log("  Проверьте адрес, порт и логин с паролем в proxy.txt.")

    def new_session(self, warm_up: bool = True) -> None:
        """Новая сессия с подменой TLS-отпечатка + прогрев на витрине WB.

        Прогрев обязателен: WB выдаёт куки (_wbauid и компанию) на основном
        домене, и запрос к card.wb.ru без них отлетает в 403.
        """
        self.close()
        proxies = self.proxy_map()
        if HAS_CURL_CFFI:
            for profile in (self.profile, "chrome"):
                try:
                    self.session = http_lib.Session(impersonate=profile, proxies=proxies)
                    break
                except Exception as exc:
                    self.log.detail(f"профиль {profile} недоступен: {exc}")
            if self.session is None:
                self.session = http_lib.Session(proxies=proxies)
        else:
            self.session = http_lib.Session()
            self.session.headers.update({"User-Agent": FALLBACK_UA})
            if proxies:
                self.session.proxies.update(proxies)

        if warm_up:
            self.warm_up()

    def warm_up(self) -> None:
        try:
            response = self.session.get(
                WB_MAIN,
                headers={
                    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                    "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
                    "Upgrade-Insecure-Requests": "1",
                },
                timeout=REQUEST_TIMEOUT,
            )
            self.warm_status = str(response.status_code)
            cookies = "; ".join(sorted(self.session.cookies.keys()))
            self.warm_cookies = cookies
            self.log.detail(f"прогрев {WB_MAIN} -> {response.status_code}, куки: {cookies or 'нет'}")
        except Exception as exc:
            self.warm_status = f"ошибка ({type(exc).__name__})"
            self.warm_cookies = ""
            self.log.detail(f"прогрев не удался: {exc}")

    def close(self) -> None:
        if self.session is not None:
            try:
                self.session.close()
            except Exception:
                pass
            self.session = None

    def rotate(self) -> None:
        """Следующий прокси и профиль браузера + свежая сессия с прогревом.

        Прокси меняем вместе с сессией, а не отдельно: куки WB привязаны к IP,
        на котором их выдали. Сменить IP и утащить старые куки — верный 403,
        поэтому после каждой смены идёт новый прогрев на витрине.
        """
        if len(self.proxies) > 1:
            self.proxy_index += 1
            self.log.detail(f"смена прокси на {mask_proxy(self.proxy)}")
        self.profile_index += 1
        self.log.detail(f"смена профиля на {self.profile}")
        self.new_session()

    # -- запросы -----------------------------------------------------------

    @staticmethod
    def api_headers() -> dict[str, str]:
        return {
            "Accept": "*/*",
            "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
            "Origin": "https://www.wildberries.ru",
            "Referer": "https://www.wildberries.ru/",
            "Sec-Fetch-Dest": "empty",
            "Sec-Fetch-Mode": "cors",
            "Sec-Fetch-Site": "cross-site",
            "x-requested-with": "XMLHttpRequest",
        }

    def build_params(self, nms: Sequence[int], endpoint: str = "", version: int = 0,
                     dest: str = "") -> dict[str, str]:
        version = version or self.version
        return {
            "appType": "1",
            "curr": "rub",
            "dest": dest or self.dest,
            "spp": "30",
            "hide_dtype": "13" if version >= 4 else "10",
            "ab_testing": "false",
            "lang": "ru",
            "nm": ";".join(str(nm) for nm in nms),
        }

    def raw_request(self, url: str, params: dict[str, str]) -> tuple[Any, str]:
        """Возвращает (payload, ошибка). payload=None, если не получилось."""
        try:
            response = self.session.get(
                url,
                params=params,
                headers=self.api_headers(),
                timeout=REQUEST_TIMEOUT,
            )
        except Exception as exc:
            # curl тащит в сообщение ссылку на свою документацию и коды —
            # в консоли это стена текста, поэтому оставляем только суть.
            message = str(exc).split(". See https://")[0].strip()
            if len(message) > 120:
                message = message[:117] + "…"
            # Отказ самого прокси легко принять за блокировку WB — разделяем.
            hint = f" (через прокси {mask_proxy(self.proxy)})" if self.proxy else ""
            return None, f"сеть: {message}{hint}"

        code = response.status_code
        if code != 200:
            body = response.text or ""
            snippet = body[:200].replace("\n", " ")
            self.log.detail(f"{url} -> {code}; тело: {snippet}")
            # В ответ на запрос API прилетела HTML-страница «доступ запрещён» —
            # это блокировка на входе, а не отказ самого API. Разница
            # принципиальная: чинить надо маршрут, а не запрос.
            lowered = body[:2000].lower()
            if code == 403 and ("<html" in lowered or "доступ" in lowered):
                return None, "HTTP 403 (страница блокировки, не ответ API)"
            return None, f"HTTP {code}"

        try:
            return response.json(), ""
        except Exception:
            text = (response.text or "").strip()
            if not text:
                return None, "пустой ответ"
            try:
                return json.loads(text), ""
            except Exception:
                return None, "ответ не JSON"

    # -- подбор рабочей комбинации ----------------------------------------

    def detect_dest(self) -> str:
        """Спрашиваем у WB код региона вместо того, чтобы угадывать."""
        payload, error = self.raw_request(
            GEO_URL,
            {"currency": "RUB", "latitude": "55.7522", "longitude": "37.6156", "locale": "ru"},
        )
        if error or not isinstance(payload, dict):
            self.log.detail(f"гео-запрос не дал dest: {error}")
            return ""
        # Настоящий код лежит в xinfo — готовой строке параметров,
        # которую WB собирает для себя: "appType=1&curr=rub&dest=...&spp=30".
        xinfo = payload.get("xinfo")
        if isinstance(xinfo, str) and xinfo:
            dest = parse_qs(xinfo).get("dest", [""])[0].strip()
            if dest:
                self.log.detail(f"dest из xinfo: {dest}")
                return dest

        for key in ("dest", "destination"):
            value = payload.get(key)
            if isinstance(value, (int, str)) and str(value).strip():
                return str(value).strip()

        # В destinations лежит несколько кодов, и нужный — не первый.
        # Актуальные коды положительные и длинные, устаревшие отрицательные.
        destinations = payload.get("destinations")
        if isinstance(destinations, list) and destinations:
            positive = [d for d in destinations if isinstance(d, int) and d > 0]
            if positive:
                return str(max(positive))
            return str(destinations[-1])
        return ""

    def probe(self, sample: Sequence[int]) -> bool:
        """Ищет комбинацию эндпоинт+dest+профиль, на которую WB отвечает 200."""
        sample = list(sample)[:5] or [1]
        self.new_session()

        dests: list[str] = []
        for candidate in (self.dest_hint, self.detect_dest(), *DEST_CANDIDATES):
            if candidate and candidate not in dests:
                dests.append(candidate)
        if self.dest_hint:
            self.log.detail(f"dest из настроек: {self.dest_hint}")

        soft_hit: tuple[str, int, str] | None = None
        rounds = min(MAX_PROBE_ROUNDS, len(IMPERSONATE_PROFILES))
        budget = MAX_PROBE_REQUESTS
        network_failures = 0
        http_failures = 0

        for attempt_round in range(rounds):
            # Со второго круга перебираем только самые вероятные регионы:
            # если дело было в dest, это выяснилось ещё на первом.
            round_dests = dests if attempt_round == 0 else dests[:2]

            for url, version in CARD_ENDPOINTS:
                for dest in round_dests:
                    if budget <= 0:
                        self.log.detail("подбор остановлен: исчерпан лимит проб")
                        return self._settle(soft_hit)
                    budget -= 1
                    time.sleep(0.25)
                    params = self.build_params(sample, url, version, dest)
                    payload, error = self.raw_request(url, params)
                    label = f"{url} dest={dest} {self.profile}"

                    if error:
                        self.attempts.append(Attempt(url, dest, self.profile, error))
                        self.log.detail(f"проба {label}: {error}")
                        if error.startswith("сеть:"):
                            network_failures += 1
                        else:
                            http_failures += 1
                        # Если до WB вообще не доходит, перебирать эндпоинты
                        # незачем: проблема не в блокировке, а в сети.
                        if http_failures == 0 and network_failures >= 6:
                            self.log.detail("подбор прекращён: сеть не отвечает")
                            return self._settle(soft_hit)
                        continue

                    products = parse_products(payload)
                    if products:
                        self.attempts.append(Attempt(url, dest, self.profile, "OK"))
                        self.endpoint, self.version, self.dest = url, version, dest
                        self.log(f"Рабочий канал: {url} (dest={dest}, профиль {self.profile})")
                        return True

                    # 200, но пусто — возможно, артикулов просто нет.
                    # Запоминаем и продолжаем искать комбинацию с товарами.
                    self.attempts.append(Attempt(url, dest, self.profile, "200, пусто"))
                    if soft_hit is None:
                        soft_hit = (url, version, dest)
                    self.log.detail(f"проба {label}: 200, но список товаров пуст")

            if attempt_round < rounds - 1:
                self.rotate()

        return self._settle(soft_hit)

    def _settle(self, soft_hit: tuple[str, int, str] | None) -> bool:
        """Товаров не нашли, но если канал отвечал 200 — работаем через него."""
        if soft_hit:
            self.endpoint, self.version, self.dest = soft_hit
            self.log(
                f"Канал отвечает, но товаров не вернул: {self.endpoint} (dest={self.dest})"
            )
            return True
        return False

    # -- рабочие запросы ---------------------------------------------------

    def shrink(self, failed_size: int) -> None:
        """Пачка не прошла — уменьшаем размер для всех следующих запросов.

        Без этого предел WB нащупывается заново на каждой пачке: лишние
        запросы, лишние блокировки и втрое дольше на тех же данных.
        """
        if failed_size >= self.batch_size and self.batch_size > MIN_BATCH:
            self.batch_size = max(MIN_BATCH, self.batch_size // 2)
            self.log(f"  уменьшаю пачку до {self.batch_size} шт")

    def fetch(self, nms: Sequence[int], attempts: int = 0) -> tuple[list[dict], str]:
        """Запрос пачки с повторами: при 403/429 меняем профиль и ждём."""
        last_error = "неизвестная ошибка"
        # Для пачки дробление помогает быстрее, чем смена профиля;
        # весь арсенал повторов тратим только на одиночный артикул.
        attempts = attempts or (MAX_ATTEMPTS if len(nms) == 1 else 2)

        for attempt in range(attempts):
            if self.delay:
                time.sleep(self.delay + random.uniform(0, 0.3))

            payload, error = self.raw_request(self.endpoint, self.build_params(nms))
            if not error:
                self.dead_streak = 0
                return parse_products(payload), ""

            last_error = error
            blocked = "403" in error or "429" in error or "401" in error
            if attempt < attempts - 1:
                if blocked:
                    self.rotate()
                    time.sleep(min(2 ** attempt, 8) + random.uniform(0, 0.5))
                else:
                    time.sleep(1.5 * (attempt + 1))

        self.dead_streak += 1
        return [], last_error


def fetch_with_split(client: WBClient, nms: Sequence[int]) -> tuple[list[dict], list[tuple[int, str]]]:
    """Не вышло пачкой — делим пополам. Одиночный отказ идёт в ошибки."""
    products, error = client.fetch(nms)
    if not error:
        return products, []

    if len(nms) > 1 and client.dead_streak < MAX_DEAD_STREAK:
        client.shrink(len(nms))
        middle = len(nms) // 2
        left_products, left_errors = fetch_with_split(client, nms[:middle])
        right_products, right_errors = fetch_with_split(client, nms[middle:])
        return left_products + right_products, left_errors + right_errors

    return [], [(nm, error) for nm in nms]


# --------------------------------------------------------------------------
# Официальный API продавца (необязательный, но не блокируется)
# --------------------------------------------------------------------------

def fetch_supplier_stocks(token: str, log: Log, proxy: str = "") -> dict[int, int]:
    """Остатки своих товаров через statistics-api. Пусто, если токена нет.

    Токен идёт внутри HTTPS, так что прокси видит только имя хоста.
    """
    if not token:
        return {}

    log("Запрашиваю остатки через API продавца…")
    try:
        proxies = {"http": proxy, "https": proxy} if proxy else None
        session = http_lib.Session(proxies=proxies) if proxies else http_lib.Session()
        response = session.get(
            STATS_URL,
            params={"dateFrom": "2019-06-20"},
            headers={"Authorization": token, "Accept": "application/json"},
            timeout=120,
        )
        if response.status_code != 200:
            log(f"! API продавца: HTTP {response.status_code} — данные не получены")
            log.detail((response.text or "")[:300])
            return {}
        rows = response.json()
    except Exception as exc:
        log(f"! API продавца недоступен: {exc}")
        return {}

    if not isinstance(rows, list):
        log("! API продавца вернул неожиданный формат")
        return {}

    stocks: dict[int, int] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        nm = row.get("nmId")
        qty = row.get("quantity", row.get("quantityFull"))
        if isinstance(nm, int) and isinstance(qty, (int, float)):
            stocks[nm] = stocks.get(nm, 0) + int(qty)

    log(f"API продавца: {len(stocks)} артикулов с остатками")
    return stocks


# --------------------------------------------------------------------------
# Отчёт
# --------------------------------------------------------------------------

@dataclass
class Row:
    nm: int
    status: str
    stock: int = 0
    source: str = ""
    name: str = ""
    brand: str = ""
    seller: str = ""
    price: float | None = None
    note: str = ""

    @property
    def url(self) -> str:
        return f"https://www.wildberries.ru/catalog/{self.nm}/detail.aspx"


COLUMNS = [
    ("Артикул", 14),
    ("Статус", 18),
    ("Остаток, шт", 13),
    ("Источник", 14),
    ("Название", 45),
    ("Бренд", 20),
    ("Продавец", 22),
    ("Цена, ₽", 12),
    ("Комментарий", 34),
    ("Ссылка", 46),
]

FILL_BY_STATUS = {
    STATUS_ON_SALE: "C6EFCE",
    STATUS_ZERO: "FFEB9C",
    STATUS_ABSENT: "FFC7CE",
    STATUS_ERROR: "D9D9D9",
}


def write_report(rows: Sequence[Row], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Остатки"

    header_fill = PatternFill("solid", fgColor="4F81BD")
    header_font = Font(bold=True, color="FFFFFF")

    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width

    for line, row in enumerate(rows, start=2):
        values = [
            row.nm,
            row.status,
            row.stock,
            row.source,
            row.name,
            row.brand,
            row.seller,
            row.price,
            row.note,
            row.url,
        ]
        for index, value in enumerate(values, start=1):
            sheet.cell(row=line, column=index, value=value)

        color = FILL_BY_STATUS.get(row.status)
        if color:
            fill = PatternFill("solid", fgColor=color)
            for index in range(1, len(COLUMNS) + 1):
                sheet.cell(row=line, column=index).fill = fill

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
    book.save(path)


# --------------------------------------------------------------------------
# Основной сценарий
# --------------------------------------------------------------------------

def build_rows(articles: Sequence[int], products: dict[int, dict],
               errors: dict[int, str], supplier: dict[int, int]) -> list[Row]:
    rows: list[Row] = []
    for nm in articles:
        product = products.get(nm)
        if product is not None:
            stock = product_stock(product)
            source = "сайт WB"
            if nm in supplier:
                stock = supplier[nm]
                source = "API продавца"
            rows.append(Row(
                nm=nm,
                status=STATUS_ON_SALE if stock > 0 else STATUS_ZERO,
                stock=stock,
                source=source,
                name=str(product.get("name") or ""),
                brand=str(product.get("brand") or ""),
                seller=str(product.get("supplier") or product.get("supplierId") or ""),
                price=product_price(product),
            ))
        elif nm in supplier:
            stock = supplier[nm]
            rows.append(Row(
                nm=nm,
                status=STATUS_ON_SALE if stock > 0 else STATUS_ZERO,
                stock=stock,
                source="API продавца",
                note="карточка не отдана сайтом",
            ))
        elif nm in errors:
            rows.append(Row(nm=nm, status=STATUS_ERROR, source="", note=errors[nm]))
        else:
            rows.append(Row(nm=nm, status=STATUS_ABSENT, source="сайт WB"))
    return rows


def print_diagnosis(client: WBClient, log: Log) -> None:
    """Когда ни одна комбинация не сработала — говорим, что именно пробовали."""
    network_only = bool(client.attempts) and all(
        attempt.result.startswith("сеть:") for attempt in client.attempts
    )

    log("")
    if network_only:
        # До WB не дошёл ни один запрос — дело не в блокировке.
        log("Соединение не устанавливается вообще. Ошибка связи:")
        log(f"  {client.attempts[0].result}")
        log("")
        if client.proxies:
            log("Так отвечает нерабочий прокси. Проверьте:")
            log("  1. Адрес, порт, логин и пароль в proxy.txt — опечатка в одном")
            log("     символе выглядит ровно так же.")
            log("  2. Оплачен ли прокси и не кончился ли срок аренды.")
            log("  3. Не привязан ли прокси к другому IP: многие продавцы пускают")
            log("     только с адреса, который вы указали в личном кабинете.")
        else:
            log("Похоже, нет доступа в интернет: фаервол, антивирус или прокси")
            log("организации. Откройте https://www.wildberries.ru в браузере —")
            log("если тоже не открывается, программа тут ни при чём.")
        log("")
        log(f"Полный протокол: {log.path}")
        return

    log("Не удалось получить ни одного ответа от WB.")
    log("")
    log(f"Прогрев www.wildberries.ru: {client.warm_status}")
    log(f"Куки после прогрева: {client.warm_cookies or 'НЕТ — это и есть причина 403'}")
    log("")
    log("Что пробовали:")
    grouped: dict[tuple[str, str, str], int] = {}
    for attempt in client.attempts:
        key = (attempt.endpoint, attempt.dest, attempt.result)
        grouped[key] = grouped.get(key, 0) + 1
    for (endpoint, dest, result), count in grouped.items():
        suffix = f" ×{count}" if count > 1 else ""
        log(f"  {endpoint} dest={dest} → {result}{suffix}")

    blocked_page = any(
        "страница блокировки" in attempt.result for attempt in client.attempts
    )

    log("")
    if blocked_page:
        # Ключевой факт: ответ пришёл от фильтра на входе, не от API.
        log("WB вернул страницу блокировки вместо ответа API.")
        log("Это значит, что запросы к card.wb.ru закрыты для вашей сети —")
        log("дело не в программе и не в артикулах.")
        log("")
        log("Проверьте за минуту: откройте на телефоне, отключив Wi-Fi,")
        log("тот же адрес card.wb.ru. Открылся текст с данными — закрыт именно")
        log("ваш домашний IP, поможет прокси (см. proxy.txt в README).")
        log("Та же страница блокировки — закрыт весь канал, и надёжнее всего")
        log("перейти на официальный API продавца по token.txt.")
        log("")
        log(f"Полный протокол: {log.path}")
        # Причина установлена точно, дальше гадать не о чем.
        return

    if client.proxies:
        log(f"Перебрано прокси: {len(client.proxies)}")
        for proxy in client.proxies:
            log(f"  {mask_proxy(proxy)}")
        log("")
        log("Вероятные причины по порядку:")
        log("  1. Прокси не российский или уже в бане у WB. Строку «Внешний IP»")
        log("     выше сверьте со страной: нужен RU, желательно резидентный")
        log("     или мобильный, а не дата-центр.")
        log("  2. Прокси живой, но WB его знает. Возьмите другой адрес")
        log("     у того же продавца — дата-центровые подсети выбиваются пачками.")
    else:
        log("Вероятные причины по порядку:")
        log("  1. Блокировка IP. VPN/прокси вне РФ, корпоративная сеть или")
        log("     хостинг — WB режет такие адреса целиком. Пропишите российский")
        log("     прокси в proxy.txt рядом с программой.")
        log("  2. Нет доступа к wildberries.ru из этой сети (фаервол, DNS).")
        log("     Откройте https://www.wildberries.ru в браузере на этом же ПК.")
    log("  3. WB снова сменил адрес API. Полный протокол запросов лежит")
    log("     в wb_stocks.log рядом с программой.")
    if not HAS_CURL_CFFI:
        log("  4. curl_cffi не установлен — подмена TLS-отпечатка не работает.")
        log("     Установите: pip install curl_cffi")
    log("")
    log("Обходной путь без публичного API: положите рядом token.txt с ключом")
    log("Statistics API продавца — остатки своих товаров тогда возьмутся оттуда.")


def run() -> int:
    setup_console()
    folder = base_dir()
    log = Log(folder / "wb_stocks.log")
    started = time.time()

    try:
        config = Config.load(folder, log)

        source = find_input_file(folder, config.input_file)
        if source is None:
            log("Не найден файл с артикулами.")
            log(f"Положите рядом с программой шаблон_артикулы.xlsx: {folder}")
            return 1

        log(f"Файл: {source}")
        articles = read_articles(source)
        if not articles:
            log("В файле не нашлось ни одного артикула (нужны числа от 5 до 12 цифр).")
            return 1

        log(f"Артикулов к проверке: {len(articles)}")

        if config.proxies:
            log(f"Прокси: {len(config.proxies)} шт — {mask_proxy(config.proxies[0])}"
                + (" и другие" if len(config.proxies) > 1 else ""))
        else:
            log("Прокси не задан — иду напрямую")

        supplier = fetch_supplier_stocks(
            config.token, log, config.proxies[0] if config.proxies else ""
        )

        client = WBClient(log, delay=config.delay, dest_hint=config.dest,
                          proxies=config.proxies)
        products: dict[int, dict] = {}
        errors: dict[int, str] = {}

        client.check_proxy()
        log("Подбираю рабочий канал к WB…")
        if client.probe(articles):
            # Значение из wb_config.json уважаем как есть: по умолчанию это 50,
            # и дальше размер всё равно подстроится сам при первом отказе.
            client.batch_size = max(MIN_BATCH, min(config.batch_size, 100))
            log(f"Опрашиваю пачками по {client.batch_size} шт")
            log("")

            queue = list(articles)
            done = 0
            while queue:
                # Размер пачки берём каждый раз заново: он мог уменьшиться.
                batch, queue = queue[:client.batch_size], queue[client.batch_size:]
                found, failed = fetch_with_split(client, batch)
                done += len(batch)

                for product in found:
                    nm = product.get("id")
                    if isinstance(nm, int):
                        products[nm] = product
                for nm, error in failed:
                    errors[nm] = error

                mark = " " if not failed else "!"
                log(f"  [{done}/{len(articles)}]{mark} карточек получено: {len(products)}"
                    + (f" — {failed[0][1]}" if failed else ""))

                if client.dead_streak >= MAX_DEAD_STREAK:
                    log("")
                    log("WB перестал отвечать — прекращаю опрос, чтобы не ждать впустую.")
                    for nm in articles:
                        if nm not in products and nm not in errors:
                            errors[nm] = "опрос прерван"
                    break
        else:
            for nm in articles:
                errors[nm] = "WB не ответил ни на одну пробу"
            print_diagnosis(client, log)

        client.close()

        rows = build_rows(articles, products, errors, supplier)
        report = folder / f"Остатки_WB_{datetime.now():%Y-%m-%d_%H-%M}.xlsx"
        write_report(rows, report)

        on_sale = sum(1 for row in rows if row.status == STATUS_ON_SALE)
        zero = sum(1 for row in rows if row.status == STATUS_ZERO)
        absent = sum(1 for row in rows if row.status == STATUS_ABSENT)
        failed_count = sum(1 for row in rows if row.status == STATUS_ERROR)
        total_stock = sum(row.stock for row in rows)

        log("")
        log("-" * 56)
        log(f"Всего артикулов: {len(rows)}")
        log(f"Найдено карточек: {on_sale + zero}")
        log(f"  в продаже: {on_sale}")
        log(f"  нулевой остаток: {zero}")
        log(f"Нет на сайте: {absent}")
        log(f"Ошибок запроса: {failed_count}")
        log(f"Суммарный остаток: {total_stock} шт")
        log("-" * 56)
        log(f"Готово за {time.time() - started:.1f} с")
        log(f"Отчёт: {report}")
        log(f"Протокол: {log.path}")

        if failed_count and failed_count < len(rows):
            log("")
            log("Часть артикулов не проверена — они помечены серым в отчёте.")
            log("Запустите ещё раз: разовые ошибки обычно уходят.")

        return 0

    except KeyboardInterrupt:
        log("")
        log("Прервано пользователем.")
        return 1
    except Exception as exc:
        import traceback

        log("")
        log(f"Сбой: {type(exc).__name__}: {exc}")
        log.detail(traceback.format_exc())
        log(f"Подробности в {log.path}")
        return 1
    finally:
        log.close()


def main() -> None:
    code = run()
    try:
        input("\nНажмите Enter, чтобы закрыть окно...")
    except (EOFError, KeyboardInterrupt):
        pass
    sys.exit(code)


if __name__ == "__main__":
    main()
